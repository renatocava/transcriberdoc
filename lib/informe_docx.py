"""Informe ecográfico en Word, escrito sobre la plantilla .docx original.

Se conserva todo lo que no es contenido (membrete, pie, márgenes, tipografía,
viñetas de la conclusión). Del documento se toma el formato de sus propios
párrafos: el de un órgano (rótulo en negrita + texto), el de una línea suelta
y el de un ítem de conclusión, y con él se escriben los párrafos nuevos.
"""

from __future__ import annotations

import copy
import datetime as dt
import io
from typing import Any

from lib.ecografias import PLANTILLAS, fecha_larga


def _negrita(run) -> bool:
    rpr = run._r.rPr
    return rpr is not None and rpr.b is not None and rpr.b.val is not False


def _rpr(run):
    return copy.deepcopy(run._r.rPr) if run._r.rPr is not None else None


def _rellenar(parrafos, rotulo: str, valor: str) -> None:
    """Pone `valor` tras los dos puntos de la línea del encabezado que empieza con `rotulo`."""
    par = next((p for p in parrafos if p.text.strip().startswith(rotulo)), None)
    if par is None:
        return
    runs = par.runs
    i = next((i for i, r in enumerate(runs) if ":" in r.text), None)
    if i is None:
        par.add_run(f"\t:  {valor}")
        return
    runs[i].text = runs[i].text[: runs[i].text.index(":")] + f":  {valor}"
    for r in runs[i + 1:]:
        r._r.getparent().remove(r._r)


class _Estilo:
    """Formato de párrafo y de texto copiado de un párrafo de la plantilla."""

    def __init__(self, parrafo, rpr_rotulo=None, rpr_texto=None):
        self.ppr = copy.deepcopy(parrafo._p.pPr) if parrafo._p.pPr is not None else None
        self.rotulo = rpr_rotulo
        self.texto = rpr_texto

    def parrafo(self, antes_de, trozos: list[tuple[Any, str]]):
        """Inserta un párrafo nuevo antes de `antes_de` con los trozos (rPr, texto)."""
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph

        p = OxmlElement("w:p")
        if self.ppr is not None:
            p.append(copy.deepcopy(self.ppr))
        antes_de._p.addprevious(p)
        par = Paragraph(p, antes_de._parent)
        for rpr, texto in trozos:
            run = par.add_run(texto)  # convierte \t en <w:tab/>
            if rpr is not None:
                run._r.insert(0, copy.deepcopy(rpr))
        return par


def _estilos_cuerpo(parrafos):
    """(órgano, línea, vacío) tomados del cuerpo de la plantilla."""
    organo = linea = vacio = None
    for p in parrafos:
        runs = [r for r in p.runs if r.text.strip()]
        if not runs:
            vacio = vacio or _Estilo(p)
            continue
        if organo is None and _negrita(runs[0]) and ":" in runs[0].text:
            cuerpo = next((r for r in runs[1:] if not _negrita(r)), runs[0])
            organo = _Estilo(p, _rpr(runs[0]), _rpr(cuerpo))
            # Sin líneas sueltas en la plantilla (transvaginal), van como el texto del órgano.
            linea_organo = _Estilo(p, None, _rpr(cuerpo))
        elif linea is None and not any(_negrita(r) for r in runs):
            linea = _Estilo(p, None, _rpr(runs[0]))
    if organo is None:
        raise ValueError("La plantilla no tiene el formato esperado (rótulo de órgano en negrita).")
    linea = linea or linea_organo
    return organo, linea, vacio or linea


def _quitar(parrafo) -> None:
    parrafo._p.getparent().remove(parrafo._p)


def generar_docx(informe: dict[str, Any]) -> bytes:
    """Bytes del .docx del informe, listo para descargar."""
    from docx import Document

    plantilla = PLANTILLAS[informe["plantilla"]]
    doc = Document(str(plantilla.ruta))
    parrafos = doc.paragraphs

    pac = informe.get("paciente") or {}
    edad = pac.get("edad")
    fecha = dt.date.fromisoformat(informe["fecha"]) if informe.get("fecha") else dt.date.today()
    _rellenar(parrafos, "NOMBRES", (pac.get("nombre") or "").upper())
    _rellenar(parrafos, "EDAD", f"{edad} AÑOS" if edad not in (None, "") else "AÑOS")
    _rellenar(parrafos, "MEDICO", (informe.get("medico") or "").upper() or "PARTICULAR")
    _rellenar(parrafos, "EXAMEN", plantilla.examen)
    _rellenar(parrafos, "FECHA", fecha_larga(fecha))

    textos = [p.text.strip() for p in parrafos]
    i_fecha = next(i for i, t in enumerate(textos) if t.startswith("FECHA"))
    i_conc = next(i for i, t in enumerate(textos) if t.startswith("CONCLUSI"))
    cuerpo = parrafos[i_fecha + 1: i_conc]
    conclusion = parrafos[i_conc]
    items = [p for p in parrafos[i_conc + 1:] if p.text.strip()]

    # Cuerpo: un bloque por órgano, separados por un párrafo vacío.
    organo, linea, vacio = _estilos_cuerpo(cuerpo)
    for p in cuerpo:
        _quitar(p)
    vacio.parrafo(conclusion, [])
    for seccion in plantilla.secciones:
        lineas = [l for l in informe["secciones"].get(seccion.clave, "").split("\n") if l.strip()]
        if seccion.titulo and lineas:
            organo.parrafo(conclusion, [(organo.rotulo, f"{seccion.titulo}: "), (organo.texto, lineas.pop(0).strip())])
        for texto in lineas:
            linea.parrafo(conclusion, [(linea.texto, texto.strip())])
        vacio.parrafo(conclusion, [])

    # Conclusión: los ítems nuevos heredan el formato (y la viñeta) del primero.
    if items:
        primero = items[0]
        con_texto = next(r for r in primero.runs if r.text.strip())
        sangria = primero.text[: len(primero.text) - len(primero.text.lstrip())]
        estilo = _Estilo(primero, None, _rpr(con_texto))
        for texto in informe.get("conclusion") or []:
            estilo.parrafo(primero, [(estilo.texto, sangria + texto.strip())])
        for p in items:
            _quitar(p)

    salida = io.BytesIO()
    doc.save(salida)
    return salida.getvalue()


def nombre_archivo(informe: dict[str, Any]) -> str:
    """«ECO ABDOMEN COMPLETO - MENDOZA HUAMAN CARLOS - 2026-10-04.docx»."""
    base = PLANTILLAS[informe["plantilla"]].archivo.removesuffix(" NORMAL.docx")
    nombre = ((informe.get("paciente") or {}).get("nombre") or "").strip().upper()
    nombre = "".join(c for c in nombre if c.isalnum() or c in " -")
    partes = [base, nombre, informe.get("fecha") or ""]
    return " - ".join(p for p in partes if p) + ".docx"
